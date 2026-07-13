import io
from unittest.mock import patch

import pytest
from docx import Document

from app.services.file_parser import (
    MAX_FILE_SIZE,
    extract_text_from_docx,
    extract_text_from_pdf,
    extract_text_from_txt,
    is_supported_file,
    parse_uploaded_file,
)


def test_extract_text_from_txt():
    text = "Hello, world!\n这是中文。"
    assert extract_text_from_txt(text.encode("utf-8")) == text


def test_extract_text_from_txt_ignores_invalid_bytes():
    raw = b"Hello \xff\xfe world"
    assert extract_text_from_txt(raw) == "Hello  world"


def test_extract_text_from_docx():
    buffer = io.BytesIO()
    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    doc.save(buffer)
    buffer.seek(0)

    result = extract_text_from_docx(buffer.read())
    assert result == "First paragraph\nSecond paragraph"


def test_extract_text_from_pdf():
    with patch("app.services.file_parser.pdfplumber.open") as mock_open:
        mock_page = mock_open.return_value.__enter__.return_value.pages.__getitem__.return_value
        mock_page.extract_text.return_value = "PDF page text"

        # Simulate single page iteration
        pages_iter = mock_open.return_value.__enter__.return_value.pages
        pages_iter.__iter__.return_value = [mock_page]

        result = extract_text_from_pdf(b"fake pdf bytes")
        assert result == "PDF page text"


def test_parse_uploaded_file_unsupported():
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_uploaded_file("image.png", b"\x89PNG\r\n\x1a\n\x00\x00\x00\x0dIHDR")


def test_parse_uploaded_file_txt():
    assert parse_uploaded_file("notes.txt", b"notes") == "notes"
    assert parse_uploaded_file("notes.md", b"# Title") == "# Title"


def test_is_supported_file():
    assert is_supported_file("file.txt")
    assert is_supported_file("file.PDF")
    assert not is_supported_file("file.exe")


def test_file_size_limit():
    oversized = b"x" * (MAX_FILE_SIZE + 1)
    with pytest.raises(ValueError, match="exceeds maximum"):
        extract_text_from_txt(oversized)
